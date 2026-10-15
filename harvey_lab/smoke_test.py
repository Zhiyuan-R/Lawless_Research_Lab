"""Keyless end-to-end smoke test for the Harvey LAB pipeline.

Runs the real `lab_core.harness.run` entry point (task loading, podman sandbox,
tool executor, agent loop, transcript/metrics) with a scripted model in place
of an LLM, then checks that the grader can locate and extract every
deliverable. No API keys are used and no judge is called, so this proves the
plumbing works, not the quality of any model.

The scripted model: globs the matter file, reads every document, writes a
placeholder for each required deliverable (.docx via python-docx, .xlsx via
openpyxl, anything else as text) from inside the sandbox, then calls finish.

Usage (from the repo root, after harvey_lab/setup.sh):
    harvey_lab/run_smoke_test.sh [task-id]
"""

import argparse
import json
import sys
from pathlib import Path

from lab_core import paths
from lab_core.evaluation.scoring import _match_deliverables, read_file_as_text
from lab_core.harness import run as harness_run
from lab_core.harness.adapters.base import ModelAdapter, ModelResponse, ToolCall

DEFAULT_TASK = "corporate-ma/review-data-room-red-flag-review"
MODEL_NAME = "scripted/smoke-test"

WRITE_DELIVERABLES = r'''
import json, sys
from pathlib import Path
names = json.loads(sys.argv[1])
for name in names:
    path = Path("/workspace/output") / name
    path.parent.mkdir(parents=True, exist_ok=True)
    text = f"SMOKE TEST PLACEHOLDER for {name}. Not legal work product."
    if name.endswith(".docx"):
        import docx
        d = docx.Document(); d.add_heading("Smoke test", 1); d.add_paragraph(text); d.save(path)
    elif name.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.Workbook(); wb.active["A1"] = text; wb.save(path)
    else:
        path.write_text(text + "\n")
    print("wrote", path)
'''


class ScriptedAdapter(ModelAdapter):
    """Plays a fixed four-turn script; message format is internal to this adapter."""

    def __init__(self, documents: list[str], deliverables: list[str]):
        super().__init__(model=MODEL_NAME)
        self.documents = documents
        self.deliverables = deliverables
        self.turn = 0

    def _calls(self, *calls: tuple[str, dict]) -> ModelResponse:
        tool_calls = [
            ToolCall(id=f"t{self.turn}-{i}", name=name, arguments=json.dumps(args))
            for i, (name, args) in enumerate(calls)
        ]
        message = {"role": "assistant", "tool_calls": [tc.__dict__ for tc in tool_calls]}
        return ModelResponse(message=message, tool_calls=tool_calls)

    def chat(self, messages: list[dict], tools: list[dict]) -> ModelResponse:
        self.turn += 1
        if self.turn == 1:
            return self._calls(("glob", {"pattern": "documents/**/*"}))
        if self.turn == 2:
            return self._calls(*[("read", {"file_path": f"documents/{d}", "limit": 40}) for d in self.documents])
        if self.turn == 3:
            script = WRITE_DELIVERABLES.replace("'", "'\\''")
            names = json.dumps(self.deliverables).replace("'", "'\\''")
            return self._calls(("bash", {"command": f"python3 -c '{script}' '{names}'"}))
        return self._calls(("finish", {"summary": "Smoke test complete.", "deliverables": self.deliverables}))

    def make_tool_result_messages(self, results: list[tuple[str, str]]) -> list[dict]:
        return [{"role": "tool", "tool_call_id": tid, "content": out} for tid, out in results]

    def make_system_message(self, content: str) -> dict:
        return {"role": "system", "content": content}

    def make_user_message(self, content: str) -> dict:
        return {"role": "user", "content": content}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("task", nargs="?", default=DEFAULT_TASK)
    cli = parser.parse_args()

    task = harness_run.load_task(cli.task)
    docs_dir = Path(task["docs_dir"])
    documents = sorted(str(p.relative_to(docs_dir)) for p in docs_dir.rglob("*") if p.is_file())
    deliverables_map = task["config"].get("deliverables", {})
    deliverables = list(deliverables_map.keys()) or ["response.md"]

    harness_run.create_adapter = lambda **_: ScriptedAdapter(documents, deliverables)
    args = harness_run.parser.parse_args(["--model", MODEL_NAME, "--task", cli.task, "--max-turns", "10"])
    harness_run.main(args)

    results_dir = paths.results_dir() / args.run_id
    metrics = json.loads((results_dir / "metrics.json").read_text())
    output_dir = results_dir / "output"
    actual = sorted(str(p.relative_to(output_dir)) for p in output_dir.rglob("*") if p.is_file())
    matched = _match_deliverables(deliverables_map, actual, output_dir)

    failures = []
    if metrics["finish_reason"] != "finish_tool":
        failures.append(f"finish_reason={metrics['finish_reason']}")
    if metrics["documents_read"] != metrics["total_documents"]:
        failures.append(f"documents_read {metrics['documents_read']}/{metrics['total_documents']}")
    print("\nGrader-side checks:")
    for name in deliverables:
        path = output_dir / name
        text = read_file_as_text(path) if path.exists() else ""
        ok = "SMOKE TEST PLACEHOLDER" in text
        print(f"  {'OK  ' if ok else 'FAIL'} {name}: extracted {len(text)} chars")
        if not ok:
            failures.append(f"deliverable {name} not extractable")
    print(f"  deliverable matching: {json.dumps(matched, default=str)[:300]}")
    if set(matched) != set(deliverables_map):
        failures.append(f"grader matched {sorted(matched)} of {sorted(deliverables_map)}")

    print("\nSMOKE TEST", "FAILED: " + "; ".join(failures) if failures else "PASSED")
    print(f"Run dir: {results_dir}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
